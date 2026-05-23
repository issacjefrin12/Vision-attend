"""Excel/CSV export service."""
from datetime import date
from io import BytesIO
from typing import List
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_

from app.models.attendance import Attendance
from app.models.student import Student
from app.models.course import Course


class ExportService:
    """Export attendance data to Excel/CSV."""
    
    async def export_attendance_excel(
        self,
        db: AsyncSession,
        course_id: int,
        start_date: date,
        end_date: date
    ) -> BytesIO:
        """
        Export attendance data to Excel with styling.
        
        Returns:
            BytesIO buffer containing Excel file
        """
        # Fetch attendance records
        result = await db.execute(
            select(Attendance, Student, Course)
            .join(Student, Attendance.student_id == Student.id)
            .join(Course, Attendance.course_id == Course.id)
            .where(
                and_(
                    Attendance.course_id == course_id,
                    Attendance.date >= start_date,
                    Attendance.date <= end_date
                )
            )
            .order_by(Attendance.date, Student.student_id)
        )
        
        records = result.fetchall()
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Report"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        present_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        late_fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
        absent_fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Headers
        headers = ["Date", "Student ID", "Student Name", "Course", "Check-in Time", "Status", "Confidence"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        
        # Data rows
        for row_idx, (attendance, student, course) in enumerate(records, 2):
            ws.cell(row=row_idx, column=1, value=str(attendance.date))
            ws.cell(row=row_idx, column=2, value=student.student_id)
            ws.cell(row=row_idx, column=3, value=student.full_name)
            ws.cell(row=row_idx, column=4, value=course.name)
            ws.cell(row=row_idx, column=5, value=attendance.check_in_time.strftime("%H:%M:%S"))
            
            status_cell = ws.cell(row=row_idx, column=6, value=attendance.status.value.upper())
            if attendance.status.value == "present":
                status_cell.fill = present_fill
            elif attendance.status.value == "late":
                status_cell.fill = late_fill
            else:
                status_cell.fill = absent_fill
            
            ws.cell(row=row_idx, column=7, value=f"{attendance.confidence_score:.2%}" if attendance.confidence_score else "N/A")
            
            # Apply border to all cells
            for col in range(1, 8):
                ws.cell(row=row_idx, column=col).border = border
        
        # Auto-adjust column widths
        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_length + 2
        
        # Save to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return buffer
    
    async def export_attendance_csv(
        self,
        db: AsyncSession,
        course_id: int,
        start_date: date,
        end_date: date
    ) -> str:
        """Export attendance data to CSV string."""
        result = await db.execute(
            select(Attendance, Student, Course)
            .join(Student, Attendance.student_id == Student.id)
            .join(Course, Attendance.course_id == Course.id)
            .where(
                and_(
                    Attendance.course_id == course_id,
                    Attendance.date >= start_date,
                    Attendance.date <= end_date
                )
            )
            .order_by(Attendance.date, Student.student_id)
        )
        
        records = result.fetchall()
        
        data = []
        for attendance, student, course in records:
            data.append({
                "Date": str(attendance.date),
                "Student ID": student.student_id,
                "Student Name": student.full_name,
                "Course": course.name,
                "Check-in Time": attendance.check_in_time.strftime("%H:%M:%S"),
                "Status": attendance.status.value,
                "Confidence": f"{attendance.confidence_score:.2%}" if attendance.confidence_score else "N/A"
            })
        
        df = pd.DataFrame(data)
        return df.to_csv(index=False)

    async def export_class_entry_attendance_excel(
        self,
        db: AsyncSession,
        class_id: int,
        target_date: "date",
    ) -> BytesIO:
        """
        Export entry attendance for a class on a given date to Excel.

        Includes summary block (Total / Present / Absent / %) at the top,
        then a per-student table with color-coded status.
        """
        from app.models.attendance import AttendanceType, AttendanceStatus
        from app.models.academic_class import AcademicClass
        from app.models.department import Department

        # Fetch class info
        cls_result = await db.execute(
            select(AcademicClass, Department)
            .join(Department, AcademicClass.department_id == Department.id)
            .where(AcademicClass.id == class_id)
        )
        cls_row = cls_result.first()
        class_label = (
            f"{cls_row[1].full_name} Year {cls_row[0].year} Sec {cls_row[0].section}"
            if cls_row
            else f"Class {class_id}"
        )

        # Fetch students in this class
        students_result = await db.execute(
            select(Student)
            .where(and_(Student.class_id == class_id, Student.is_active == True))
            .order_by(Student.register_no)
        )
        students = list(students_result.scalars().all())

        # Fetch entry attendance for this class on the target date
        attendance_result = await db.execute(
            select(Attendance)
            .join(Student, Attendance.student_id == Student.id)
            .where(
                and_(
                    Student.class_id == class_id,
                    Attendance.date == target_date,
                    Attendance.attendance_type == AttendanceType.ENTRY.value,
                    Attendance.status.in_(
                        [AttendanceStatus.PRESENT, AttendanceStatus.LATE]
                    ),
                )
            )
        )
        attendance_records = attendance_result.scalars().all()

        # Build lookup: student_id -> attendance record
        attendance_map: dict[int, "Attendance"] = {}
        for record in attendance_records:
            if record.student_id not in attendance_map:
                attendance_map[record.student_id] = record

        total_students = len(students)
        present_count = sum(1 for s in students if s.id in attendance_map)
        absent_count = total_students - present_count
        percentage = round(
            (present_count / total_students * 100) if total_students > 0 else 0.0, 2
        )

        # ── Build workbook ──────────────────────────────────────────
        wb = Workbook()
        ws = wb.active
        ws.title = "Entry Attendance"

        # Styles
        title_font = Font(bold=True, size=14, color="FFFFFF")
        label_font = Font(bold=True, size=11, color="FFFFFF")
        value_font = Font(size=11, color="FFFFFF")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        summary_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
        present_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        late_fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
        absent_fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        center = Alignment(horizontal="center", vertical="center")

        # Row 1 — Title
        ws.merge_cells("A1:F1")
        title_cell = ws.cell(row=1, column=1, value=f"Entry Attendance — {class_label}")
        title_cell.font = title_font
        title_cell.fill = summary_fill
        title_cell.alignment = center

        # Row 2 — Date
        ws.merge_cells("A2:F2")
        date_cell = ws.cell(row=2, column=1, value=f"Date: {target_date}")
        date_cell.font = label_font
        date_cell.fill = summary_fill
        date_cell.alignment = center

        # Row 3 — Summary stats
        summary_labels = ["Total Students", "Present", "Absent", "Percentage"]
        summary_values = [total_students, present_count, absent_count, f"{percentage}%"]
        for col_idx, (label, value) in enumerate(
            zip(summary_labels, summary_values), 1
        ):
            lbl_cell = ws.cell(row=3, column=col_idx, value=label)
            lbl_cell.font = label_font
            lbl_cell.fill = summary_fill
            lbl_cell.alignment = center

            val_cell = ws.cell(row=4, column=col_idx, value=value)
            val_cell.font = value_font
            val_cell.fill = summary_fill
            val_cell.alignment = center

        # Row 5 — spacer (blank)

        # Row 6 — Table header
        headers = ["Register No", "Name", "Status", "Type", "Date", "Time"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        # Row 7+ — Student data
        for row_idx, student in enumerate(students, 7):
            record = attendance_map.get(student.id)

            ws.cell(row=row_idx, column=1, value=student.register_no).border = border
            ws.cell(row=row_idx, column=2, value=student.full_name).border = border

            status_text = "Present" if record else "Absent"
            if record and record.status.value == "late":
                status_text = "Late"
            status_cell = ws.cell(row=row_idx, column=3, value=status_text)
            status_cell.border = border
            status_cell.alignment = center
            if record:
                if record.status.value == "late":
                    status_cell.fill = late_fill
                else:
                    status_cell.fill = present_fill
                status_cell.font = Font(bold=True, color="FFFFFF")
            else:
                status_cell.fill = absent_fill
                status_cell.font = Font(bold=True, color="FFFFFF")

            type_cell = ws.cell(
                row=row_idx,
                column=4,
                value=record.attendance_type if record else "-",
            )
            type_cell.border = border
            type_cell.alignment = center

            date_cell = ws.cell(
                row=row_idx,
                column=5,
                value=str(record.date) if record else "-",
            )
            date_cell.border = border
            date_cell.alignment = center

            time_cell = ws.cell(
                row=row_idx,
                column=6,
                value=(
                    record.check_in_time.strftime("%I:%M %p")
                    if record and record.check_in_time
                    else "-"
                ),
            )
            time_cell.border = border
            time_cell.alignment = center

        # Auto-adjust column widths
        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(max_length + 3, 12)

        # Save to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer


export_service = ExportService()
